#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
rank_candidates_v4_layered.py

V4 layered morphology-adaptive ranking layer for contact_field_v3 output.

Input:
    contact_field_v3_*.csv

Output:
    One Excel workbook with multiple sheets:
      01_raw_original          original v3 data, unsorted
      02_scored_all            original order + v4 scores + filter labels
      03_filter_broad          broad-filter-passed candidates, unsorted
      04_filter_medium         medium-filter-passed candidates, unsorted
      05_filter_strict         strict-filter-passed candidates, unsorted
      06_ranked_final_all      all valid candidates ranked by final_score_v4
      07_ranked_final_broad    broad pool ranked by final_score_v4
      08_ranked_final_medium   medium pool ranked by final_score_v4
      09_ranked_final_strict   strict pool ranked by final_score_v4
      10_rank_line             linear-target mode ranking
      11_rank_plane            plane-target mode ranking
      12_rank_compact          compact-target mode ranking
      13_rank_roughness        contact-field roughness/smoothness ranking
      14_rank_microfit         microdomain-fit ranking
      15_references            reference candidates, if provided
      16_weights_filters       target morphology, dynamic weights, filter thresholds
      17_metric_notes          metric explanations

Also writes a TXT report explaining:
    - target morphology
    - dynamic mode weights
    - broad / medium / strict filter rules
    - scoring weights
    - top candidates
    - reference candidate ranks

Important:
    v3 = feature extractor from PDB.
    v4 = ranking/scoring layer from v3 CSV.

    The reference candidates are only sanity checks. They are NOT used to tune weights.

Author: generated with ChatGPT, based on the user's RFdiffusion binder-ranking project.
"""

import argparse
import csv
import math
import sys
from pathlib import Path
from typing import Dict, List, Tuple, Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
except Exception:
    print("ERROR: openpyxl is required to write multi-sheet Excel output.", file=sys.stderr)
    print("Install it in the binderranker conda env:", file=sys.stderr)
    print("  pip install -i https://pypi.tuna.tsinghua.edu.cn/simple openpyxl", file=sys.stderr)
    sys.exit(1)


EPS = 1e-12


# ============================================================
# Basic helpers
# ============================================================

def safe_float(x: Any, default: float = math.nan) -> float:
    if x is None:
        return default
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip()
    if s == "" or s.lower() in {"nan", "none", "null"}:
        return default
    try:
        return float(s)
    except Exception:
        return default


def is_valid_row(row: Dict[str, Any]) -> bool:
    e = str(row.get("error", "")).strip()
    return e == "" or e.lower() in {"nan", "none", "null"}


def finite_values(rows: List[Dict[str, Any]], col: str) -> List[float]:
    vals = []
    for r in rows:
        v = safe_float(r.get(col))
        if math.isfinite(v):
            vals.append(v)
    return vals


def quantile(vals: List[float], q: float) -> float:
    vals = sorted([v for v in vals if math.isfinite(v)])
    if not vals:
        return math.nan
    q = max(0.0, min(1.0, q))
    if len(vals) == 1:
        return vals[0]
    pos = q * (len(vals) - 1)
    lo = int(math.floor(pos))
    hi = int(math.ceil(pos))
    if lo == hi:
        return vals[lo]
    frac = pos - lo
    return vals[lo] * (1.0 - frac) + vals[hi] * frac


def median(vals: List[float]) -> float:
    return quantile(vals, 0.5)


def clip01(x: float) -> float:
    if not math.isfinite(x):
        return 0.5
    return max(0.0, min(1.0, x))


def robust_norm_value(x: float, lo: float, hi: float, higher_better: bool = True) -> float:
    """
    Robust percentile normalization.
    Values below q_low -> 0, above q_high -> 1.
    If lower is better, reverse after normalization.
    """
    if not math.isfinite(x) or not math.isfinite(lo) or not math.isfinite(hi) or abs(hi - lo) < EPS:
        return 0.5
    z = (x - lo) / (hi - lo)
    z = clip01(z)
    return z if higher_better else 1.0 - z


def normalize_weights(raw: Dict[str, float]) -> Dict[str, float]:
    s = sum(max(0.0, v) for v in raw.values())
    if s <= EPS:
        return {k: 0.0 for k in raw}
    return {k: max(0.0, v) / s for k, v in raw.items()}


# ============================================================
# Derived metrics
# ============================================================

def add_derived_metrics(rows: List[Dict[str, Any]]) -> None:
    """
    Add derived metrics from v3 CSV columns.
    Does not touch PDB files.
    """
    for r in rows:
        raw8 = safe_float(r.get("raw_mean_shell_8"), 0.0)
        raw12 = safe_float(r.get("raw_mean_shell_12"), 0.0)

        if raw12 > EPS:
            r["shell_sensitivity_12_vs_8"] = (raw12 - raw8) / (raw12 + EPS)
            r["tight_contact_fraction_8_over_12"] = raw8 / (raw12 + EPS)
        else:
            r["shell_sensitivity_12_vs_8"] = 1.0
            r["tight_contact_fraction_8_over_12"] = 0.0

        eff = safe_float(r.get("effective_weight_sum"), 0.0)
        t_cov = safe_float(r.get("target_effective_coverage"), 0.0)
        t_len = safe_float(r.get("target_len"), math.nan)
        t_neff_norm = safe_float(r.get("target_contact_neff_norm"), math.nan)

        if math.isfinite(t_len) and math.isfinite(t_neff_norm):
            t_neff = max(t_neff_norm * t_len, EPS)
            r["contact_density_target_neff"] = eff / t_neff
        else:
            r["contact_density_target_neff"] = math.nan

        if math.isfinite(t_len):
            covered_count = max(t_cov * t_len, EPS)
            r["contact_density_target_covered"] = eff / covered_count
        else:
            r["contact_density_target_covered"] = math.nan

        center_dist = safe_float(r.get("micro_center_distance_weighted"), math.nan)
        if math.isfinite(center_dist):
            # v3 default ideal_dist=7.5, micro_dist_sigma=3.0
            r["micro_center_distance_score"] = math.exp(-((center_dist - 7.5) ** 2) / (2.0 * 3.0 * 3.0))
        else:
            r["micro_center_distance_score"] = math.nan

        r["contacted_linearity_score"] = safe_float(r.get("contacted_target_linearity"), math.nan)
        r["contacted_planarity_score"] = safe_float(r.get("contacted_target_planarity"), math.nan)
        r["contacted_compactness_score"] = safe_float(r.get("contacted_target_compactness"), math.nan)


def infer_target_morphology(rows: List[Dict[str, Any]], alpha: float = 1.25, min_mode_weight: float = 0.05) -> Dict[str, float]:
    """
    Infer target-level morphology from median target_linearity/planarity/compactness.

    For one RFdiffusion job, target is fixed, so median across valid rows should be stable.
    Mode weights are softened to avoid completely shutting off plane/compact scores.
    """
    valid = [r for r in rows if is_valid_row(r)]

    L = median(finite_values(valid, "target_linearity"))
    P = median(finite_values(valid, "target_planarity"))
    C = median(finite_values(valid, "target_compactness"))

    if not math.isfinite(L): L = 0.0
    if not math.isfinite(P): P = 0.0
    if not math.isfinite(C): C = 0.0

    raw = {
        "line": max(L, 0.0) ** alpha,
        "plane": max(P, 0.0) ** alpha,
        "compact": max(C, 0.0) ** alpha,
    }
    base = normalize_weights(raw)

    floor = max(0.0, min(min_mode_weight, 0.30))
    if 3.0 * floor >= 1.0:
        floor = 0.0

    mode = {}
    for k, v in base.items():
        mode[k] = floor + (1.0 - 3.0 * floor) * v
    mode = normalize_weights(mode)

    return {
        "target_linearity_median": L,
        "target_planarity_median": P,
        "target_compactness_median": C,
        "morph_alpha": alpha,
        "min_mode_weight": min_mode_weight,
        "mode_weight_line": mode["line"],
        "mode_weight_plane": mode["plane"],
        "mode_weight_compact": mode["compact"],
    }


# ============================================================
# Scoring components
# ============================================================

def get_component_weights() -> Dict[str, Dict[str, Tuple[str, float, bool]]]:
    """
    component_name -> (column_name, weight, higher_better)
    Weights are normalized inside each score.
    """
    score_line = {
        # For elongated / helical-bundle targets.
        "effective_contact_strength": ("effective_weight_sum", 0.17, True),
        "target_coverage": ("target_effective_coverage", 0.19, True),
        "target_span": ("target_contact_span_norm", 0.19, True),
        "target_neff": ("target_contact_neff_norm", 0.08, True),
        "contact_map_continuity": ("contact_map_continuity_score", 0.09, True),
        "binder_active_smooth": ("binder_field_active_smoothness", 0.08, True),
        "contacted_target_linearity": ("contacted_linearity_score", 0.06, True),
        "cb_closer": ("cb_closer_weight_ratio", 0.05, True),
        "front_facing": ("front_facing_weight_ratio", 0.04, True),
        "tight_contact_fraction": ("tight_contact_fraction_8_over_12", 0.04, True),
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.07, False),
        "low_contact_jump": ("contact_map_jump_fraction", 0.04, False),
        "low_shell_sensitivity": ("shell_sensitivity_12_vs_8", 0.03, False),
    }

    score_plane = {
        # For flatter target surfaces.
        "micro_fit": ("micro_fit_score_weighted", 0.20, True),
        "micro_plane_alignment": ("micro_plane_alignment_weighted", 0.14, True),
        "micro_planarity": ("micro_planarity_mean", 0.10, True),
        "micro_center_distance": ("micro_center_distance_score", 0.08, True),
        "target_coverage": ("target_effective_coverage", 0.13, True),
        "effective_contact_strength": ("effective_weight_sum", 0.12, True),
        "contacted_target_planarity": ("contacted_planarity_score", 0.08, True),
        "contact_map_continuity": ("contact_map_continuity_score", 0.05, True),
        "cb_closer": ("cb_closer_weight_ratio", 0.04, True),
        "front_facing": ("front_facing_weight_ratio", 0.03, True),
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.07, False),
        "low_contact_jump": ("contact_map_jump_fraction", 0.03, False),
    }

    score_compact = {
        # For compact/pocket-like target patches.
        "effective_contact_strength": ("effective_weight_sum", 0.18, True),
        "contact_density_neff": ("contact_density_target_neff", 0.15, True),
        "target_coverage": ("target_effective_coverage", 0.10, True),
        "micro_fit": ("micro_fit_score_weighted", 0.10, True),
        "contacted_target_compactness": ("contacted_compactness_score", 0.07, True),
        "cb_closer": ("cb_closer_weight_ratio", 0.08, True),
        "front_facing": ("front_facing_weight_ratio", 0.06, True),
        "binder_core_smooth": ("binder_field_core_smoothness", 0.07, True),
        "binder_active_smooth": ("binder_field_active_smoothness", 0.04, True),
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.12, False),
        "low_contact_jump": ("contact_map_jump_fraction", 0.04, False),
        "low_shell_sensitivity": ("shell_sensitivity_12_vs_8", 0.04, False),
    }

    score_roughness = {
        # Standalone score for the user's contact-field roughness idea.
        "binder_active_smooth": ("binder_field_active_smoothness", 0.32, True),
        "binder_core_smooth": ("binder_field_core_smoothness", 0.24, True),
        "target_active_smooth": ("target_field_active_smoothness", 0.12, True),
        "target_core_smooth": ("target_field_core_smoothness", 0.08, True),
        "contact_map_continuity": ("contact_map_continuity_score", 0.14, True),
        "low_contact_jump": ("contact_map_jump_fraction", 0.05, False),
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.05, False),
    }

    score_microfit = {
        # Standalone score for the user's microdomain-surface-complementarity idea.
        "micro_fit": ("micro_fit_score_weighted", 0.38, True),
        "micro_plane_alignment": ("micro_plane_alignment_weighted", 0.22, True),
        "micro_planarity": ("micro_planarity_mean", 0.14, True),
        "micro_center_distance": ("micro_center_distance_score", 0.10, True),
        "effective_contact_strength": ("effective_weight_sum", 0.05, True),
        "target_coverage": ("target_effective_coverage", 0.05, True),
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.04, False),
        "low_contact_jump": ("contact_map_jump_fraction", 0.02, False),
    }

    score_safety = {
        # Stabilizer, not the main objective.
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.24, False),
        "cb_closer": ("cb_closer_weight_ratio", 0.16, True),
        "low_contact_jump": ("contact_map_jump_fraction", 0.16, False),
        "binder_active_smooth": ("binder_field_active_smoothness", 0.14, True),
        "tight_contact_fraction": ("tight_contact_fraction_8_over_12", 0.10, True),
        "low_shell_sensitivity": ("shell_sensitivity_12_vs_8", 0.10, False),
        "front_facing": ("front_facing_weight_ratio", 0.06, True),
        "low_clash_pairs": ("clash_pairs", 0.04, False),
    }

    return {
        "score_line": score_line,
        "score_plane": score_plane,
        "score_compact": score_compact,
        "score_roughness": score_roughness,
        "score_microfit": score_microfit,
        "score_safety": score_safety,
    }


def collect_metric_specs(component_weights: Dict[str, Dict[str, Tuple[str, float, bool]]]) -> Dict[Tuple[str, bool], None]:
    specs = {}
    for comps in component_weights.values():
        for _, (col, _, higher) in comps.items():
            specs[(col, higher)] = None
    return specs


def compute_normalization_ranges(rows: List[Dict[str, Any]], specs: Dict[Tuple[str, bool], None],
                                 q_low: float, q_high: float) -> Dict[Tuple[str, bool], Tuple[float, float]]:
    valid = [r for r in rows if is_valid_row(r)]
    ranges = {}
    for col, higher in specs:
        vals = finite_values(valid, col)
        ranges[(col, higher)] = (quantile(vals, q_low), quantile(vals, q_high))
    return ranges


def add_weighted_score(row_index: int,
                       norm_cache: Dict[Tuple[int, str, bool], float],
                       components: Dict[str, Tuple[str, float, bool]]) -> float:
    total_w = sum(max(0.0, w) for _, w, _ in components.values())
    if total_w <= EPS:
        return 0.0
    score = 0.0
    for _, (col, w, higher) in components.items():
        score += (max(0.0, w) / total_w) * norm_cache.get((row_index, col, higher), 0.5)
    return clip01(score)


def compute_scores(rows: List[Dict[str, Any]], morph: Dict[str, float],
                   q_low: float = 0.05, q_high: float = 0.95) -> Dict[str, Any]:
    comp_weights = get_component_weights()
    specs = collect_metric_specs(comp_weights)
    ranges = compute_normalization_ranges(rows, specs, q_low, q_high)

    norm_cache: Dict[Tuple[int, str, bool], float] = {}
    for i, r in enumerate(rows):
        for (col, higher), (lo, hi) in ranges.items():
            val = safe_float(r.get(col))
            norm_cache[(i, col, higher)] = robust_norm_value(val, lo, hi, higher_better=higher)

    for i, r in enumerate(rows):
        if not is_valid_row(r):
            for score_name in list(comp_weights.keys()) + ["morphology_adaptive_score", "final_score_v4", "rank_consensus_score"]:
                r[score_name] = math.nan
            continue

        for score_name, comps in comp_weights.items():
            r[score_name] = add_weighted_score(i, norm_cache, comps)

        adaptive = (
            morph["mode_weight_line"] * safe_float(r.get("score_line"), 0.0)
            + morph["mode_weight_plane"] * safe_float(r.get("score_plane"), 0.0)
            + morph["mode_weight_compact"] * safe_float(r.get("score_compact"), 0.0)
        )
        r["morphology_adaptive_score"] = clip01(adaptive)

        r["final_score_v4"] = clip01(
            0.85 * safe_float(r.get("morphology_adaptive_score"), 0.0)
            + 0.10 * safe_float(r.get("score_safety"), 0.0)
            + 0.05 * safe_float(r.get("score_roughness"), 0.0)
        )

    add_rank_consensus(rows, morph)

    return {
        "component_weights": comp_weights,
        "normalization_ranges": ranges,
        "q_low": q_low,
        "q_high": q_high,
    }


def add_rank_consensus(rows: List[Dict[str, Any]], morph: Dict[str, float]) -> None:
    """
    Robust rank-aggregation score from multiple mode scores.
    """
    valid = [r for r in rows if is_valid_row(r)]
    score_cols = [
        ("score_line", morph["mode_weight_line"]),
        ("score_plane", morph["mode_weight_plane"]),
        ("score_compact", morph["mode_weight_compact"]),
        ("score_roughness", 0.10),
        ("score_microfit", 0.05),
        ("score_safety", 0.10),
    ]
    total = sum(w for _, w in score_cols)
    score_cols = [(c, w / total) for c, w in score_cols]

    percentiles = {}
    for col, _ in score_cols:
        sorted_rows = sorted(valid, key=lambda r: safe_float(r.get(col), -1e99))
        n = len(sorted_rows)
        p = {}
        for idx, r in enumerate(sorted_rows):
            p[str(r.get("pdb_name", ""))] = idx / max(n - 1, 1)
        percentiles[col] = p

    for r in rows:
        if not is_valid_row(r):
            r["rank_consensus_score"] = math.nan
            continue
        name = str(r.get("pdb_name", ""))
        s = 0.0
        for col, w in score_cols:
            s += w * percentiles[col].get(name, 0.0)
        r["rank_consensus_score"] = clip01(s)


# ============================================================
# Layered filters
# ============================================================

def filter_definition() -> Dict[str, Dict[str, Tuple[str, str, float]]]:
    """
    Three filter layers.

    Format:
      layer -> rule_name -> (column, mode, quantile)
    mode:
      "min": require value >= quantile(column, q)
      "max": require value <= quantile(column, q)

    Quantiles are deliberately percentile-based, so the script adapts to each job.
    """
    return {
        "broad": {
            "effective_weight_sum": ("effective_weight_sum", "min", 0.10),
            "target_effective_coverage": ("target_effective_coverage", "min", 0.10),
            "target_contact_span_norm": ("target_contact_span_norm", "min", 0.10),
            "backfacing_cb_far_weight_ratio": ("backfacing_cb_far_weight_ratio", "max", 0.95),
            "binder_field_active_roughness": ("binder_field_active_roughness", "max", 0.95),
            "contact_map_jump_fraction": ("contact_map_jump_fraction", "max", 0.95),
            "shell_sensitivity_12_vs_8": ("shell_sensitivity_12_vs_8", "max", 0.98),
            "clash_pairs": ("clash_pairs", "max", 0.99),
        },
        "medium": {
            # Main contact/coverage requirements become meaningful but not extreme.
            "effective_weight_sum": ("effective_weight_sum", "min", 0.40),
            "target_effective_coverage": ("target_effective_coverage", "min", 0.40),
            "target_contact_span_norm": ("target_contact_span_norm", "min", 0.40),

            # Direction/contact safety.
            "backfacing_cb_far_weight_ratio": ("backfacing_cb_far_weight_ratio", "max", 0.85),
            "binder_field_active_roughness": ("binder_field_active_roughness", "max", 0.85),
            "contact_map_jump_fraction": ("contact_map_jump_fraction", "max", 0.85),
            "shell_sensitivity_12_vs_8": ("shell_sensitivity_12_vs_8", "max", 0.90),
            "cb_closer_weight_ratio": ("cb_closer_weight_ratio", "min", 0.20),
            "contact_map_continuity_score": ("contact_map_continuity_score", "min", 0.20),
            "clash_pairs": ("clash_pairs", "max", 0.99),
        },
        "strict": {
            # Strict is a candidate-pool filter, not final truth.
            # It should keep a reasonable elite pool rather than only a few structures.
            "effective_weight_sum": ("effective_weight_sum", "min", 0.60),
            "target_effective_coverage": ("target_effective_coverage", "min", 0.60),
            "target_contact_span_norm": ("target_contact_span_norm", "min", 0.55),

            # Risk controls, still not too harsh.
            "backfacing_cb_far_weight_ratio": ("backfacing_cb_far_weight_ratio", "max", 0.75),
            "binder_field_active_roughness": ("binder_field_active_roughness", "max", 0.75),
            "contact_map_jump_fraction": ("contact_map_jump_fraction", "max", 0.75),
            "shell_sensitivity_12_vs_8": ("shell_sensitivity_12_vs_8", "max", 0.80),
            "cb_closer_weight_ratio": ("cb_closer_weight_ratio", "min", 0.35),
            "contact_map_continuity_score": ("contact_map_continuity_score", "min", 0.35),
            "score_safety": ("score_safety", "min", 0.35),
            "clash_pairs": ("clash_pairs", "max", 0.99),
        },
    }


def compute_filter_thresholds(rows: List[Dict[str, Any]]) -> Dict[str, Dict[str, Dict[str, Any]]]:
    valid = [r for r in rows if is_valid_row(r)]
    definitions = filter_definition()
    thresholds: Dict[str, Dict[str, Dict[str, Any]]] = {}

    for layer, rules in definitions.items():
        thresholds[layer] = {}
        for rule_name, (col, mode, q) in rules.items():
            vals = finite_values(valid, col)
            thr = quantile(vals, q)
            if col == "clash_pairs":
                # If all are zero, keep zero. If rare non-zero, 99th percentile may still be 0.
                thr = max(0.0, thr)
            thresholds[layer][rule_name] = {
                "column": col,
                "mode": mode,
                "quantile": q,
                "threshold": thr,
            }
    return thresholds


def apply_layered_filters(rows: List[Dict[str, Any]], thresholds: Dict[str, Dict[str, Dict[str, Any]]]) -> None:
    for r in rows:
        if not is_valid_row(r):
            for layer in ["broad", "medium", "strict"]:
                r[f"filter_{layer}_pass"] = "NO"
                r[f"filter_{layer}_reasons"] = "error_not_empty"
            r["filter_level"] = "FAIL"
            continue

        passed_layers = []

        for layer in ["broad", "medium", "strict"]:
            reasons = []
            for rule_name, spec in thresholds[layer].items():
                col = spec["column"]
                mode = spec["mode"]
                thr = spec["threshold"]
                val = safe_float(r.get(col))
                if not math.isfinite(val) or not math.isfinite(thr):
                    continue
                if mode == "min" and val < thr:
                    reasons.append(f"low_{col}")
                elif mode == "max" and val > thr:
                    reasons.append(f"high_{col}")

            r[f"filter_{layer}_pass"] = "YES" if not reasons else "NO"
            r[f"filter_{layer}_reasons"] = ";".join(reasons)
            if not reasons:
                passed_layers.append(layer)

        if r["filter_strict_pass"] == "YES":
            r["filter_level"] = "STRICT"
        elif r["filter_medium_pass"] == "YES":
            r["filter_level"] = "MEDIUM"
        elif r["filter_broad_pass"] == "YES":
            r["filter_level"] = "BROAD"
        else:
            r["filter_level"] = "FAIL"


# ============================================================
# Ranking and references
# ============================================================

def add_ranks(rows: List[Dict[str, Any]], score_cols: List[str]) -> None:
    valid = [r for r in rows if is_valid_row(r)]

    for col in score_cols:
        sorted_rows = sorted(valid, key=lambda r: safe_float(r.get(col), -1e99), reverse=True)
        for idx, r in enumerate(sorted_rows, start=1):
            r[f"rank_{col}"] = idx
        for r in rows:
            if not is_valid_row(r):
                r[f"rank_{col}"] = ""

    # Filter-specific final ranks.
    for layer in ["broad", "medium", "strict"]:
        passed = [r for r in valid if r.get(f"filter_{layer}_pass") == "YES"]
        sorted_passed = sorted(passed, key=lambda r: safe_float(r.get("final_score_v4"), -1e99), reverse=True)
        for idx, r in enumerate(sorted_passed, start=1):
            r[f"rank_final_in_{layer}"] = idx
        for r in rows:
            if r.get(f"filter_{layer}_pass") != "YES":
                r[f"rank_final_in_{layer}"] = ""


def names_match(row_name: str, ref: str) -> bool:
    row_name = str(row_name).strip()
    ref = str(ref).strip().replace(".pdb", "")
    if not row_name or not ref:
        return False
    return row_name == ref or row_name.endswith(ref)


def find_references(rows: List[Dict[str, Any]], refs: List[str]) -> List[Dict[str, Any]]:
    out = []
    for ref in refs:
        found = None
        for r in rows:
            if names_match(str(r.get("pdb_name", "")), ref):
                found = r
                break
        if found is None:
            out.append({
                "_reference_query": ref,
                "pdb_name": "NOT_FOUND",
                "filter_level": "",
                "filter_broad_pass": "",
                "filter_medium_pass": "",
                "filter_strict_pass": "",
                "filter_strict_reasons": "reference_not_found",
            })
        else:
            item = dict(found)
            item["_reference_query"] = ref
            out.append(item)
    return out


# ============================================================
# Excel output
# ============================================================

DISPLAY_COLS = [
    "pdb_name",
    "filter_level",
    "filter_broad_pass", "filter_broad_reasons",
    "filter_medium_pass", "filter_medium_reasons",
    "filter_strict_pass", "filter_strict_reasons",
    "final_score_v4",
    "morphology_adaptive_score",
    "rank_consensus_score",
    "score_line",
    "score_plane",
    "score_compact",
    "score_roughness",
    "score_microfit",
    "score_safety",
    "rank_final_score_v4",
    "rank_final_in_broad",
    "rank_final_in_medium",
    "rank_final_in_strict",
    "rank_morphology_adaptive_score",
    "rank_rank_consensus_score",
    "effective_weight_sum",
    "target_effective_coverage",
    "target_contact_span_norm",
    "target_contact_neff_norm",
    "target_contact_entropy",
    "raw_mean_shell_8",
    "raw_mean_shell_10",
    "raw_mean_shell_12",
    "tight_contact_fraction_8_over_12",
    "shell_sensitivity_12_vs_8",
    "front_facing_weight_ratio",
    "tangential_weight_ratio",
    "backfacing_cb_far_weight_ratio",
    "cb_closer_weight_ratio",
    "cb_delta_weighted_mean",
    "binder_field_active_roughness",
    "binder_field_active_smoothness",
    "binder_field_core_roughness",
    "binder_field_core_smoothness",
    "target_field_active_roughness",
    "target_field_active_smoothness",
    "target_field_core_roughness",
    "target_field_core_smoothness",
    "contact_map_continuity_score",
    "contact_map_jump_fraction",
    "micro_plane_alignment_weighted",
    "micro_center_distance_weighted",
    "micro_center_distance_score",
    "micro_planarity_mean",
    "micro_fit_score_weighted",
    "target_linearity",
    "target_planarity",
    "target_compactness",
    "contacted_target_linearity",
    "contacted_target_planarity",
    "contacted_target_compactness",
    "desired_weight_fraction",
    "undesired_weight_fraction",
    "clash_pairs",
    "min_distance_global",
    "error",
]


def fmt(v: Any) -> Any:
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return ""
    return v


def autosize_sheet(ws, max_width: int = 42) -> None:
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col[:200]:
            try:
                max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), max_width)


def write_sheet(ws, rows: List[Dict[str, Any]], headers: List[str]) -> None:
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([fmt(r.get(h, "")) for h in headers])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    autosize_sheet(ws)


def write_kv_sheet(ws, title: str, rows: List[Tuple[Any, Any, Any]]) -> None:
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["key", "value", "note"])
    for cell in ws[3]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A4"
    autosize_sheet(ws)


def metric_notes() -> List[Dict[str, str]]:
    rows = [
        ("final_score_v4", "Primary score", "0.85 adaptive + 0.10 safety + 0.05 roughness"),
        ("morphology_adaptive_score", "Dynamic score", "line/plane/compact mixed by target morphology"),
        ("score_line", "Linear-target mode", "For elongated targets; emphasizes coverage, span, continuity, active smoothness"),
        ("score_plane", "Plane-target mode", "For flat surfaces; emphasizes microdomain plane complementarity"),
        ("score_compact", "Compact-target mode", "For pockets/local patches; emphasizes dense and safe local contact"),
        ("score_roughness", "Contact-field roughness score", "Standalone score for contact smoothness idea"),
        ("score_microfit", "Microdomain fit score", "Standalone score for microdomain surface complementarity idea"),
        ("score_safety", "Safety stabilizer", "Backfacing, jumps, shell sensitivity, clash-related stabilizer"),
        ("filter_broad_pass", "Broad filter", "Loose outlier removal; should not be used alone as final candidate pool"),
        ("filter_medium_pass", "Medium filter", "Moderate candidate pool"),
        ("filter_strict_pass", "Strict filter", "Elite candidate pool; still not ground truth"),
        ("target_effective_coverage", "Target coverage", "Fraction of target residues effectively contacted"),
        ("target_contact_span_norm", "Target span", "Spatial spread of contacted target region normalized by full target span"),
        ("effective_weight_sum", "Effective contact amount", "Sum of v3 soft contact matrix W"),
        ("backfacing_cb_far_weight_ratio", "Suspicious direction", "CA contact but pseudo-Cβ faces away and is farther"),
        ("binder_field_active_smoothness", "Boundary smoothness", "exp(-active roughness); higher means contact boundary smoother"),
        ("binder_field_core_smoothness", "Core smoothness", "exp(-core roughness); higher means contact core smoother"),
        ("contact_map_continuity_score", "Contact-map continuity", "Adjacent binder residues map to nearby target regions"),
        ("contact_map_jump_fraction", "Contact-map jump", "Fraction of active edges mapping to distant target positions"),
        ("micro_fit_score_weighted", "Microdomain fit", "local plane alignment × distance score × planarity"),
        ("shell_sensitivity_12_vs_8", "Shell sensitivity", "(raw12-raw8)/raw12; high means loose-shell contact dominates"),
        ("tight_contact_fraction_8_over_12", "Tight-contact support", "raw8/raw12; high means contact is not only loose 12A contact"),
    ]
    return [{"metric": a, "meaning": b, "note": c} for a, b, c in rows]


def build_weight_filter_rows(morph: Dict[str, float],
                             thresholds: Dict[str, Dict[str, Dict[str, Any]]],
                             metadata: Dict[str, Any]) -> List[Tuple[Any, Any, Any]]:
    out = []
    out.append(("TARGET_MORPHOLOGY", "", ""))
    out.append(("target_linearity_median", morph["target_linearity_median"], "median across valid rows"))
    out.append(("target_planarity_median", morph["target_planarity_median"], "median across valid rows"))
    out.append(("target_compactness_median", morph["target_compactness_median"], "median across valid rows"))
    out.append(("morph_alpha", morph["morph_alpha"], "power used to convert morphology into mode weights"))
    out.append(("min_mode_weight", morph["min_mode_weight"], "minimum retained mode weight"))
    out.append(("mode_weight_line", morph["mode_weight_line"], "line-target mode weight"))
    out.append(("mode_weight_plane", morph["mode_weight_plane"], "plane-target mode weight"))
    out.append(("mode_weight_compact", morph["mode_weight_compact"], "compact-target mode weight"))
    out.append(("", "", ""))

    out.append(("FINAL_SCORE", "0.85*morphology_adaptive + 0.10*safety + 0.05*roughness", "primary score"))
    out.append(("", "", ""))

    out.append(("FILTER_THRESHOLDS", "", "percentile-based thresholds"))
    for layer in ["broad", "medium", "strict"]:
        out.append((f"{layer.upper()}_FILTER", "", ""))
        for rule_name, spec in thresholds[layer].items():
            out.append((
                f"{layer}.{rule_name}",
                spec["threshold"],
                f"{spec['mode']} {spec['column']} at q={spec['quantile']}"
            ))
    out.append(("", "", ""))

    out.append(("COMPONENT_WEIGHTS", "", "weights inside each score; normalized within score"))
    comp_weights = metadata["component_weights"]
    for score_name, comps in comp_weights.items():
        total = sum(w for _, w, _ in comps.values())
        out.append((score_name, "", ""))
        for comp_name, (col, w, higher) in comps.items():
            direction = "higher_better" if higher else "lower_better"
            out.append((f"{score_name}.{comp_name}", w / total if total > EPS else w, f"{col}; {direction}"))

    out.append(("", "", ""))
    out.append(("NORMALIZATION", "", f"robust percentile q={metadata['q_low']} to q={metadata['q_high']}"))
    for (col, higher), (lo, hi) in metadata["normalization_ranges"].items():
        direction = "higher_better" if higher else "lower_better"
        out.append((f"range.{col}.{direction}", f"{lo} to {hi}", ""))
    return out


def write_excel(output_xlsx: Path,
                raw_original: List[Dict[str, Any]],
                rows: List[Dict[str, Any]],
                metadata: Dict[str, Any],
                morph: Dict[str, float],
                thresholds: Dict[str, Dict[str, Dict[str, Any]]],
                references: List[Dict[str, Any]],
                top_n_sheets: int = 500) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    raw_headers = list(raw_original[0].keys()) if raw_original else []

    headers = []
    for h in DISPLAY_COLS:
        if h in rows[0] or h in raw_headers:
            headers.append(h)
    for h in list(rows[0].keys()):
        if h not in headers:
            headers.append(h)

    ws = wb.create_sheet("01_raw_original")
    write_sheet(ws, raw_original, raw_headers)

    ws = wb.create_sheet("02_scored_all")
    write_sheet(ws, rows, headers)

    layer_specs = [
        ("03_filter_broad", "broad"),
        ("04_filter_medium", "medium"),
        ("05_filter_strict", "strict"),
    ]
    for sheet, layer in layer_specs:
        subset = [r for r in rows if r.get(f"filter_{layer}_pass") == "YES"]
        ws = wb.create_sheet(sheet)
        write_sheet(ws, subset, headers)

    ranking_sheets = [
        ("06_ranked_final_all", None, "final_score_v4"),
        ("07_ranked_final_broad", "broad", "final_score_v4"),
        ("08_ranked_final_medium", "medium", "final_score_v4"),
        ("09_ranked_final_strict", "strict", "final_score_v4"),
        ("10_rank_line", None, "score_line"),
        ("11_rank_plane", None, "score_plane"),
        ("12_rank_compact", None, "score_compact"),
        ("13_rank_roughness", None, "score_roughness"),
        ("14_rank_microfit", None, "score_microfit"),
    ]
    valid_rows = [r for r in rows if is_valid_row(r)]
    for sheet, layer, score_col in ranking_sheets:
        if layer:
            subset = [r for r in valid_rows if r.get(f"filter_{layer}_pass") == "YES"]
        else:
            subset = valid_rows
        ranked = sorted(subset, key=lambda r: safe_float(r.get(score_col), -1e99), reverse=True)
        ws = wb.create_sheet(sheet)
        write_sheet(ws, ranked[:top_n_sheets], headers)

    ws = wb.create_sheet("15_references")
    ref_headers = ["_reference_query"] + headers
    write_sheet(ws, references, ref_headers)

    ws = wb.create_sheet("16_weights_filters")
    write_kv_sheet(ws, "V4 layered morphology-adaptive ranking weights and filters",
                   build_weight_filter_rows(morph, thresholds, metadata))

    ws = wb.create_sheet("17_metric_notes")
    write_sheet(ws, metric_notes(), ["metric", "meaning", "note"])

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)


# ============================================================
# TXT report
# ============================================================

def write_txt_report(output_txt: Path,
                     rows: List[Dict[str, Any]],
                     morph: Dict[str, float],
                     thresholds: Dict[str, Dict[str, Dict[str, Any]]],
                     references: List[Dict[str, Any]],
                     top_k: int = 30) -> None:
    valid = [r for r in rows if is_valid_row(r)]
    pools = {
        "all": valid,
        "broad": [r for r in valid if r.get("filter_broad_pass") == "YES"],
        "medium": [r for r in valid if r.get("filter_medium_pass") == "YES"],
        "strict": [r for r in valid if r.get("filter_strict_pass") == "YES"],
    }

    lines = []
    lines.append("V4 layered morphology-adaptive ranking report")
    lines.append("=" * 70)
    lines.append("")
    lines.append("Target morphology:")
    lines.append(f"  target_linearity_median   = {morph['target_linearity_median']:.6f}")
    lines.append(f"  target_planarity_median   = {morph['target_planarity_median']:.6f}")
    lines.append(f"  target_compactness_median = {morph['target_compactness_median']:.6f}")
    lines.append("")
    lines.append("Dynamic mode weights:")
    lines.append(f"  line    = {morph['mode_weight_line']:.6f}")
    lines.append(f"  plane   = {morph['mode_weight_plane']:.6f}")
    lines.append(f"  compact = {morph['mode_weight_compact']:.6f}")
    lines.append("")
    lines.append("Primary score:")
    lines.append("  final_score_v4 = 0.85*morphology_adaptive_score + 0.10*score_safety + 0.05*score_roughness")
    lines.append("")
    lines.append("Pool sizes:")
    for k, v in pools.items():
        lines.append(f"  {k:>6}: {len(v)}")
    lines.append("")

    lines.append("Filter thresholds:")
    for layer in ["broad", "medium", "strict"]:
        lines.append(f"  [{layer}]")
        for rule_name, spec in thresholds[layer].items():
            lines.append(f"    {rule_name}: {spec['mode']} {spec['column']} {spec['threshold']}  (q={spec['quantile']})")
    lines.append("")

    cols = [
        "pdb_name", "filter_level", "final_score_v4",
        "morphology_adaptive_score", "score_line", "score_plane", "score_compact",
        "score_roughness", "score_microfit", "score_safety",
        "effective_weight_sum", "target_effective_coverage", "target_contact_span_norm",
        "backfacing_cb_far_weight_ratio", "binder_field_active_smoothness",
        "micro_fit_score_weighted",
    ]

    for pool_name in ["all", "broad", "medium", "strict"]:
        ranked = sorted(pools[pool_name], key=lambda r: safe_float(r.get("final_score_v4"), -1e99), reverse=True)
        lines.append(f"Top {min(top_k, len(ranked))} in {pool_name} pool by final_score_v4:")
        for idx, r in enumerate(ranked[:top_k], start=1):
            parts = []
            for c in cols:
                v = r.get(c, "")
                if isinstance(v, float):
                    parts.append(f"{c}={v:.5f}")
                else:
                    parts.append(f"{c}={v}")
            lines.append(f"  #{idx}: " + " | ".join(parts))
        lines.append("")

    lines.append("Reference candidates:")
    if references:
        for r in references:
            q = r.get("_reference_query", "")
            name = r.get("pdb_name", "NOT_FOUND")
            if name == "NOT_FOUND":
                lines.append(f"  {q}: NOT_FOUND")
            else:
                lines.append(
                    f"  {q} -> {name}: "
                    f"rank_all={r.get('rank_final_score_v4','')}, "
                    f"rank_broad={r.get('rank_final_in_broad','')}, "
                    f"rank_medium={r.get('rank_final_in_medium','')}, "
                    f"rank_strict={r.get('rank_final_in_strict','')}, "
                    f"level={r.get('filter_level','')}, "
                    f"final={safe_float(r.get('final_score_v4')):.5f}, "
                    f"line={safe_float(r.get('score_line')):.5f}, "
                    f"rough={safe_float(r.get('score_roughness')):.5f}, "
                    f"micro={safe_float(r.get('score_microfit')):.5f}, "
                    f"strict_reasons={r.get('filter_strict_reasons','')}"
                )
    else:
        lines.append("  None")
    lines.append("")
    lines.append("Reminder:")
    lines.append("  - Broad / medium / strict are candidate pools, not ground truth.")
    lines.append("  - Strict pool is intentionally meaningful but not absolute; a candidate outside strict can still be valuable.")
    lines.append("  - Reference candidates are not used to tune weights.")
    lines.append("  - MPNN + AF2 is still required for downstream validation.")
    lines.append("")

    output_txt.parent.mkdir(parents=True, exist_ok=True)
    output_txt.write_text("\n".join(lines), encoding="utf-8")


# ============================================================
# Main
# ============================================================

def read_csv_rows(path: Path) -> List[Dict[str, Any]]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return [dict(r) for r in reader]


def main() -> None:
    parser = argparse.ArgumentParser(description="V4 layered morphology-adaptive ranking for contact_field_v3 CSV.")
    parser.add_argument("--input_csv", required=True, help="Input contact_field_v3 CSV")
    parser.add_argument("--output_xlsx", required=True, help="Output Excel workbook")
    parser.add_argument("--output_txt", default=None, help="Output TXT report")
    parser.add_argument("--reference_names", default="QbQc_219_PDs_4978,QbQc_219_PDs_3722",
                        help="Comma-separated references, e.g. QbQc_219_PDs_4978,QbQc_219_PDs_3722")
    parser.add_argument("--top_n_sheets", type=int, default=500, help="Max rows in each ranked sheet")
    parser.add_argument("--top_k_report", type=int, default=30, help="Top K shown in TXT report")
    parser.add_argument("--morph_alpha", type=float, default=1.25, help="Power for morphology-to-mode conversion")
    parser.add_argument("--min_mode_weight", type=float, default=0.05, help="Minimum weight retained for each mode")
    parser.add_argument("--q_low", type=float, default=0.05, help="Low quantile for robust normalization")
    parser.add_argument("--q_high", type=float, default=0.95, help="High quantile for robust normalization")
    args = parser.parse_args()

    input_csv = Path(args.input_csv)
    output_xlsx = Path(args.output_xlsx)
    output_txt = Path(args.output_txt) if args.output_txt else output_xlsx.with_suffix(".txt")

    rows = read_csv_rows(input_csv)
    if not rows:
        raise SystemExit("No rows found in input CSV.")

    raw_original = [dict(r) for r in rows]

    add_derived_metrics(rows)
    morph = infer_target_morphology(rows, alpha=args.morph_alpha, min_mode_weight=args.min_mode_weight)
    metadata = compute_scores(rows, morph, q_low=args.q_low, q_high=args.q_high)

    thresholds = compute_filter_thresholds(rows)
    apply_layered_filters(rows, thresholds)

    rank_cols = [
        "final_score_v4",
        "morphology_adaptive_score",
        "rank_consensus_score",
        "score_line",
        "score_plane",
        "score_compact",
        "score_roughness",
        "score_microfit",
        "score_safety",
    ]
    add_ranks(rows, rank_cols)

    refs = [x.strip() for x in str(args.reference_names).split(",") if x.strip()]
    ref_rows = find_references(rows, refs)

    write_excel(output_xlsx, raw_original, rows, metadata, morph, thresholds, ref_rows, top_n_sheets=args.top_n_sheets)
    write_txt_report(output_txt, rows, morph, thresholds, ref_rows, top_k=args.top_k_report)

    valid = [r for r in rows if is_valid_row(r)]
    broad = [r for r in valid if r.get("filter_broad_pass") == "YES"]
    medium = [r for r in valid if r.get("filter_medium_pass") == "YES"]
    strict = [r for r in valid if r.get("filter_strict_pass") == "YES"]

    print(f"Wrote Excel workbook: {output_xlsx}")
    print(f"Wrote TXT report: {output_txt}")
    print("")
    print("Target morphology:")
    print(f"  linearity   = {morph['target_linearity_median']:.6f}")
    print(f"  planarity   = {morph['target_planarity_median']:.6f}")
    print(f"  compactness = {morph['target_compactness_median']:.6f}")
    print("Mode weights:")
    print(f"  line    = {morph['mode_weight_line']:.6f}")
    print(f"  plane   = {morph['mode_weight_plane']:.6f}")
    print(f"  compact = {morph['mode_weight_compact']:.6f}")
    print("")
    print("Pool sizes:")
    print(f"  all    = {len(valid)}")
    print(f"  broad  = {len(broad)}")
    print(f"  medium = {len(medium)}")
    print(f"  strict = {len(strict)}")

    ranked = sorted(strict if strict else medium if medium else broad if broad else valid,
                    key=lambda r: safe_float(r.get("final_score_v4"), -1e99),
                    reverse=True)
    print("")
    print("Top candidates in best available pool:")
    for idx, r in enumerate(ranked[:10], start=1):
        print(
            f"  #{idx:02d} {r.get('pdb_name')} "
            f"level={r.get('filter_level')} "
            f"final={safe_float(r.get('final_score_v4')):.5f} "
            f"line={safe_float(r.get('score_line')):.5f} "
            f"coverage={safe_float(r.get('target_effective_coverage')):.5f} "
            f"span={safe_float(r.get('target_contact_span_norm')):.5f}"
        )


if __name__ == "__main__":
    main()
