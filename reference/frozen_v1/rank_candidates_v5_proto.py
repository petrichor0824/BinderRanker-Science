#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
rank_candidates_v5_proto.py

BinderRanker v5 prototype ranking layer.

Design goals
------------
1. DO NOT modify frozen v4.
2. Reuse the exact frozen v4 implementation for all existing v4 component scores
   and layered filters.
3. Add candidate-specific local contacted-patch morphology and blend it with the
   v4 global target morphology.
4. Add explicit, separately inspectable interface-extent variants instead of
   silently changing the old score.
5. Keep every new architectural choice as a CLI parameter so calibration can be
   performed later on development targets without touching held-out labels.
6. Default settings are V4-COMPATIBLE: final_score_v5 should reproduce
   final_score_v4 (up to floating-point noise).

This is a DEVELOPMENT prototype, not a frozen production model.

Input
-----
contact_field_v3 CSV (same input used by rank_candidates_v4_layered.py)

Outputs
-------
- full candidate TSV with v4 + v5 fields
- multi-sheet XLSX for inspection
- TXT report with configuration and top candidates

Scientific idea
---------------
Global target morphology tells us what the whole target looks like.  The binder,
however, only sees a contacted patch.  V5 therefore constructs candidate-specific
local mode weights from contacted_target_linearity / planarity / compactness and
blends those weights with the global v4 weights:

    w_blend = (1-beta) * w_global + beta * w_local

The local vector can be shrunk toward the neutral (1/3,1/3,1/3) vector using
local_lambda and optionally calibrated with plane/compact gains.

V5 also exposes several extent variants based on normalized coverage coupled to
one geometric/support signal using a geometric mean.  No variant is assumed to
be correct a priori; the choice is an explicit parameter for later calibration.
"""

import argparse
import csv
import hashlib
import importlib.util
import math
import sys
from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


FROZEN_V4_SHA256 = "9871434c132ceb3fb0f5e54ff8ada17db241525e7f03c2124b6c6e746558ddfe"
EPS = 1e-12


# -----------------------------------------------------------------------------
# Frozen-v4 loading
# -----------------------------------------------------------------------------

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def load_frozen_v4(path: Path):
    if not path.is_file():
        raise SystemExit(f"[FATAL] frozen v4 script not found: {path}")

    got = sha256_file(path)
    if got != FROZEN_V4_SHA256:
        raise SystemExit(
            "[FATAL] frozen v4 SHA256 mismatch.\n"
            f"  expected: {FROZEN_V4_SHA256}\n"
            f"  got     : {got}\n"
            f"  file    : {path}\n"
            "Refusing to run because v5-proto comparisons must use the frozen v4 baseline."
        )

    spec = importlib.util.spec_from_file_location("binderranker_frozen_v4", path)
    if spec is None or spec.loader is None:
        raise SystemExit(f"[FATAL] cannot import frozen v4 module: {path}")

    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module, got


# -----------------------------------------------------------------------------
# Small helpers
# -----------------------------------------------------------------------------

def clip01(x: float) -> float:
    if not math.isfinite(x):
        return math.nan
    return max(0.0, min(1.0, x))


def normalize3(line: float, plane: float, compact: float) -> Tuple[float, float, float]:
    vals = [max(0.0, line), max(0.0, plane), max(0.0, compact)]
    s = sum(vals)
    if s <= EPS:
        return (1.0 / 3.0, 1.0 / 3.0, 1.0 / 3.0)
    return tuple(v / s for v in vals)


def floor_mode_weights(weights: Tuple[float, float, float], floor: float) -> Tuple[float, float, float]:
    floor = max(0.0, min(float(floor), 0.30))
    if 3.0 * floor >= 1.0:
        floor = 0.0
    out = tuple(floor + (1.0 - 3.0 * floor) * w for w in weights)
    return normalize3(*out)


def local_mode_weights(
    row: Dict[str, Any],
    v4,
    alpha: float,
    floor: float,
    local_lambda: float,
    gain_plane: float,
    gain_compact: float,
) -> Tuple[float, float, float]:
    """Candidate-specific line/plane/compact weights from contacted target patch."""
    L = v4.safe_float(row.get("contacted_linearity_score"), math.nan)
    P = v4.safe_float(row.get("contacted_planarity_score"), math.nan)
    C = v4.safe_float(row.get("contacted_compactness_score"), math.nan)

    if not math.isfinite(L):
        L = 0.0
    if not math.isfinite(P):
        P = 0.0
    if not math.isfinite(C):
        C = 0.0

    # Gains are calibration knobs, not fitted here.
    raw_L = max(L, 0.0) ** alpha
    raw_P = max(P * gain_plane, 0.0) ** alpha
    raw_C = max(C * gain_compact, 0.0) ** alpha
    base = normalize3(raw_L, raw_P, raw_C)

    lam = max(0.0, min(1.0, local_lambda))
    neutral = 1.0 / 3.0
    shrunk = tuple(neutral + lam * (w - neutral) for w in base)
    return floor_mode_weights(shrunk, floor)


def geometric_mean(values: List[float]) -> float:
    vals = [clip01(v) for v in values]
    if any(not math.isfinite(v) for v in vals):
        return math.nan
    if not vals:
        return math.nan
    product = 1.0
    for v in vals:
        product *= max(0.0, v)
    return product ** (1.0 / len(vals))


def robust_norm_column(rows: List[Dict[str, Any]], col: str, v4, q_low: float, q_high: float) -> List[float]:
    valid_vals = []
    for r in rows:
        if not v4.is_valid_row(r):
            continue
        x = v4.safe_float(r.get(col), math.nan)
        if math.isfinite(x):
            valid_vals.append(x)

    lo = v4.quantile(valid_vals, q_low)
    hi = v4.quantile(valid_vals, q_high)

    out = []
    for r in rows:
        x = v4.safe_float(r.get(col), math.nan)
        if not v4.is_valid_row(r) or not math.isfinite(x):
            out.append(math.nan)
        else:
            out.append(v4.robust_norm_value(x, lo, hi, higher_better=True))
    return out


# -----------------------------------------------------------------------------
# V5 augmentation
# -----------------------------------------------------------------------------

def add_v5_scores(
    rows: List[Dict[str, Any]],
    morph: Dict[str, float],
    v4,
    beta_local: float,
    local_lambda: float,
    gain_plane: float,
    gain_compact: float,
    morph_alpha: float,
    min_mode_weight: float,
    q_low: float,
    q_high: float,
    extent_variant: str,
    micro_weight: float,
    extent_weight: float,
    safety_weight: float,
    roughness_weight: float,
) -> Dict[str, Any]:
    beta = max(0.0, min(1.0, beta_local))

    fixed = micro_weight + extent_weight + safety_weight + roughness_weight
    if min(micro_weight, extent_weight, safety_weight, roughness_weight) < -EPS:
        raise SystemExit("[FATAL] v5 outer weights must be non-negative")
    if fixed > 1.0 + EPS:
        raise SystemExit(
            f"[FATAL] v5 fixed outer weights sum to {fixed:.6f} > 1.0"
        )
    morphology_weight = max(0.0, 1.0 - fixed)

    # Normalize a few candidate-level interface signals within each target.
    coverage_n = robust_norm_column(rows, "target_effective_coverage", v4, q_low, q_high)
    cb_support_n = robust_norm_column(rows, "cb_supported_weight_ratio", v4, q_low, q_high)
    micro_raw_n = robust_norm_column(rows, "micro_fit_score_weighted", v4, q_low, q_high)
    continuity_n = robust_norm_column(rows, "contact_map_continuity_score", v4, q_low, q_high)

    global_w = (
        float(morph["mode_weight_line"]),
        float(morph["mode_weight_plane"]),
        float(morph["mode_weight_compact"]),
    )

    for i, row in enumerate(rows):
        if not v4.is_valid_row(row):
            for col in [
                "v5_local_weight_line", "v5_local_weight_plane", "v5_local_weight_compact",
                "v5_blend_weight_line", "v5_blend_weight_plane", "v5_blend_weight_compact",
                "morphology_local_score_v5", "morphology_blended_score_v5",
                "v5_norm_coverage", "v5_norm_cb_support", "v5_norm_microfit_raw",
                "v5_norm_continuity", "extent_coverage_v5", "extent_coverage_cb_v5",
                "extent_coverage_micro_v5", "extent_coverage_continuity_v5",
                "extent_coverage_cb_micro_v5", "extent_score_v5", "final_score_v5",
            ]:
                row[col] = math.nan
            continue

        local_w = local_mode_weights(
            row=row,
            v4=v4,
            alpha=morph_alpha,
            floor=min_mode_weight,
            local_lambda=local_lambda,
            gain_plane=gain_plane,
            gain_compact=gain_compact,
        )

        blend_w = normalize3(
            (1.0 - beta) * global_w[0] + beta * local_w[0],
            (1.0 - beta) * global_w[1] + beta * local_w[1],
            (1.0 - beta) * global_w[2] + beta * local_w[2],
        )

        row["v5_local_weight_line"] = local_w[0]
        row["v5_local_weight_plane"] = local_w[1]
        row["v5_local_weight_compact"] = local_w[2]
        row["v5_blend_weight_line"] = blend_w[0]
        row["v5_blend_weight_plane"] = blend_w[1]
        row["v5_blend_weight_compact"] = blend_w[2]

        s_line = v4.safe_float(row.get("score_line"), 0.0)
        s_plane = v4.safe_float(row.get("score_plane"), 0.0)
        s_compact = v4.safe_float(row.get("score_compact"), 0.0)

        row["morphology_local_score_v5"] = clip01(
            local_w[0] * s_line + local_w[1] * s_plane + local_w[2] * s_compact
        )
        row["morphology_blended_score_v5"] = clip01(
            blend_w[0] * s_line + blend_w[1] * s_plane + blend_w[2] * s_compact
        )

        row["v5_norm_coverage"] = coverage_n[i]
        row["v5_norm_cb_support"] = cb_support_n[i]
        row["v5_norm_microfit_raw"] = micro_raw_n[i]
        row["v5_norm_continuity"] = continuity_n[i]

        row["extent_coverage_v5"] = coverage_n[i]
        row["extent_coverage_cb_v5"] = geometric_mean([coverage_n[i], cb_support_n[i]])
        row["extent_coverage_micro_v5"] = geometric_mean([coverage_n[i], micro_raw_n[i]])
        row["extent_coverage_continuity_v5"] = geometric_mean([coverage_n[i], continuity_n[i]])
        row["extent_coverage_cb_micro_v5"] = geometric_mean(
            [coverage_n[i], cb_support_n[i], micro_raw_n[i]]
        )

        variant_map = {
            "coverage": "extent_coverage_v5",
            "coverage_cb": "extent_coverage_cb_v5",
            "coverage_micro": "extent_coverage_micro_v5",
            "coverage_continuity": "extent_coverage_continuity_v5",
            "coverage_cb_micro": "extent_coverage_cb_micro_v5",
        }
        extent_col = variant_map[extent_variant]
        row["extent_score_v5"] = v4.safe_float(row.get(extent_col), 0.0)

        row["final_score_v5"] = clip01(
            morphology_weight * v4.safe_float(row.get("morphology_blended_score_v5"), 0.0)
            + micro_weight * v4.safe_float(row.get("score_microfit"), 0.0)
            + extent_weight * v4.safe_float(row.get("extent_score_v5"), 0.0)
            + safety_weight * v4.safe_float(row.get("score_safety"), 0.0)
            + roughness_weight * v4.safe_float(row.get("score_roughness"), 0.0)
        )

    return {
        "beta_local": beta,
        "local_lambda": local_lambda,
        "gain_plane": gain_plane,
        "gain_compact": gain_compact,
        "extent_variant": extent_variant,
        "morphology_weight": morphology_weight,
        "micro_weight": micro_weight,
        "extent_weight": extent_weight,
        "safety_weight": safety_weight,
        "roughness_weight": roughness_weight,
        "q_low": q_low,
        "q_high": q_high,
    }


# -----------------------------------------------------------------------------
# Ranking / output
# -----------------------------------------------------------------------------

def add_v5_ranks(rows: List[Dict[str, Any]], v4) -> None:
    valid = [r for r in rows if v4.is_valid_row(r)]

    score_cols = [
        "final_score_v5",
        "morphology_local_score_v5",
        "morphology_blended_score_v5",
        "extent_score_v5",
    ]

    for col in score_cols:
        ranked = sorted(
            valid,
            key=lambda r: (
                -v4.safe_float(r.get(col), -1e99),
                str(r.get("pdb_name", "")),
            ),
        )
        for idx, r in enumerate(ranked, start=1):
            r[f"rank_{col}"] = idx

    for layer in ["broad", "medium", "strict"]:
        passed = [r for r in valid if r.get(f"filter_{layer}_pass") == "YES"]
        ranked = sorted(
            passed,
            key=lambda r: (
                -v4.safe_float(r.get("final_score_v5"), -1e99),
                str(r.get("pdb_name", "")),
            ),
        )
        for idx, r in enumerate(ranked, start=1):
            r[f"rank_final_v5_in_{layer}"] = idx
        for r in rows:
            if r.get(f"filter_{layer}_pass") != "YES":
                r[f"rank_final_v5_in_{layer}"] = ""


def ordered_headers(rows: List[Dict[str, Any]]) -> List[str]:
    preferred = [
        "pdb_name",
        "filter_level",
        "filter_broad_pass", "filter_broad_reasons",
        "filter_medium_pass", "filter_medium_reasons",
        "filter_strict_pass", "filter_strict_reasons",
        "final_score_v5", "rank_final_score_v5",
        "rank_final_v5_in_broad", "rank_final_v5_in_medium", "rank_final_v5_in_strict",
        "morphology_blended_score_v5", "morphology_local_score_v5",
        "v5_blend_weight_line", "v5_blend_weight_plane", "v5_blend_weight_compact",
        "v5_local_weight_line", "v5_local_weight_plane", "v5_local_weight_compact",
        "extent_score_v5", "extent_coverage_v5", "extent_coverage_cb_v5",
        "extent_coverage_micro_v5", "extent_coverage_continuity_v5",
        "extent_coverage_cb_micro_v5",
        "v5_norm_coverage", "v5_norm_cb_support", "v5_norm_microfit_raw", "v5_norm_continuity",
        "final_score_v4", "rank_final_score_v4", "morphology_adaptive_score",
        "score_line", "score_plane", "score_compact", "score_microfit", "score_safety", "score_roughness",
        "target_effective_coverage", "target_contact_span_norm", "target_contact_neff_norm",
        "cb_supported_weight_ratio", "micro_fit_score_weighted", "contact_map_continuity_score",
        "target_linearity", "target_planarity", "target_compactness",
        "contacted_target_linearity", "contacted_target_planarity", "contacted_target_compactness",
        "error",
    ]

    present = list(rows[0].keys()) if rows else []
    out = []
    for h in preferred:
        if h in present and h not in out:
            out.append(h)
    for h in present:
        if h not in out:
            out.append(h)
    return out


def write_tsv(path: Path, rows: List[Dict[str, Any]], headers: List[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=headers, delimiter="\t", extrasaction="ignore")
        w.writeheader()
        for row in rows:
            clean = {}
            for h in headers:
                v = row.get(h, "")
                if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
                    v = ""
                clean[h] = v
            w.writerow(clean)


def autosize(ws, max_width: int = 42) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 0
        for cell in col:
            if cell.value is not None:
                width = max(width, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(width + 2, 9), max_width)


def write_table_sheet(ws, rows: List[Dict[str, Any]], headers: List[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        vals = []
        for h in headers:
            v = row.get(h, "")
            if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
                v = ""
            vals.append(v)
        ws.append(vals)
    ws.freeze_panes = "A2"
    autosize(ws)


def write_config_sheet(ws, config: Dict[str, Any], morph: Dict[str, float], v4_sha: str) -> None:
    ws.append(["key", "value", "note"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    rows = [
        ("model", "BinderRanker v5 prototype", "development only; not frozen"),
        ("frozen_v4_sha256", v4_sha, "must match historical baseline"),
        ("global_mode_weight_line", morph["mode_weight_line"], "v4 whole-target morphology"),
        ("global_mode_weight_plane", morph["mode_weight_plane"], "v4 whole-target morphology"),
        ("global_mode_weight_compact", morph["mode_weight_compact"], "v4 whole-target morphology"),
    ]
    for k, v in config.items():
        rows.append((k, v, "v5 prototype parameter"))

    rows += [
        ("formula", "wmorph*morph_blend + wmicro*microfit + wextent*extent + wsafety*safety + wrough*roughness", "weights sum to 1"),
        ("filters", "unchanged from frozen v4", "kept fixed in v5 prototype"),
    ]
    for row in rows:
        ws.append(list(row))
    autosize(ws)


def write_excel(
    path: Path,
    raw_original: List[Dict[str, Any]],
    rows: List[Dict[str, Any]],
    headers: List[str],
    config: Dict[str, Any],
    morph: Dict[str, float],
    v4_sha: str,
    top_n: int,
    v4,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    raw_headers = list(raw_original[0].keys()) if raw_original else []
    ws = wb.create_sheet("01_raw_original")
    write_table_sheet(ws, raw_original, raw_headers)

    ws = wb.create_sheet("02_scored_all")
    write_table_sheet(ws, rows, headers)

    for idx, layer in enumerate(["broad", "medium", "strict"], start=3):
        subset = [r for r in rows if r.get(f"filter_{layer}_pass") == "YES"]
        ws = wb.create_sheet(f"{idx:02d}_filter_{layer}")
        write_table_sheet(ws, subset, headers)

    valid = [r for r in rows if v4.is_valid_row(r)]
    specs = [
        ("06_ranked_v5_all", None, "final_score_v5"),
        ("07_ranked_v5_broad", "broad", "final_score_v5"),
        ("08_ranked_v5_medium", "medium", "final_score_v5"),
        ("09_ranked_v5_strict", "strict", "final_score_v5"),
        ("10_ranked_v4_all", None, "final_score_v4"),
        ("11_rank_v5_morph", None, "morphology_blended_score_v5"),
        ("12_rank_v5_extent", None, "extent_score_v5"),
        ("13_rank_v4_microfit", None, "score_microfit"),
    ]

    for name, layer, col in specs:
        subset = valid if layer is None else [r for r in valid if r.get(f"filter_{layer}_pass") == "YES"]
        ranked = sorted(
            subset,
            key=lambda r: (-v4.safe_float(r.get(col), -1e99), str(r.get("pdb_name", ""))),
        )
        ws = wb.create_sheet(name)
        write_table_sheet(ws, ranked[:top_n], headers)

    ws = wb.create_sheet("14_config")
    write_config_sheet(ws, config, morph, v4_sha)

    ws = wb.create_sheet("15_notes")
    notes = [
        {"item": "status", "note": "Development prototype; do not treat as frozen v5."},
        {"item": "v4 preservation", "note": "All v4 component scores and filters are computed by the SHA-guarded frozen v4 script."},
        {"item": "local morphology", "note": "Candidate-specific contacted patch chooses line/plane/compact emphasis, then blends with global target morphology."},
        {"item": "extent variants", "note": "Coverage may be used alone or geometrically coupled to CB support, microfit, or continuity. Variant choice is explicit."},
        {"item": "hyperparameters", "note": "beta/local_lambda/gains/outer weights are development parameters; choose later using target-level calibration."},
        {"item": "heldout", "note": "Do not use Core10 held-out AF2 labels until the model and hyperparameters are frozen."},
    ]
    write_table_sheet(ws, notes, ["item", "note"])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_report(
    path: Path,
    rows: List[Dict[str, Any]],
    config: Dict[str, Any],
    morph: Dict[str, float],
    v4_sha: str,
    top_k: int,
    v4,
) -> None:
    valid = [r for r in rows if v4.is_valid_row(r)]
    pools = {
        "all": valid,
        "broad": [r for r in valid if r.get("filter_broad_pass") == "YES"],
        "medium": [r for r in valid if r.get("filter_medium_pass") == "YES"],
        "strict": [r for r in valid if r.get("filter_strict_pass") == "YES"],
    }

    lines = []
    lines.append("BinderRanker v5 prototype report")
    lines.append("=" * 76)
    lines.append("DEVELOPMENT ONLY -- NOT A FROZEN MODEL")
    lines.append("")
    lines.append(f"Frozen v4 SHA256: {v4_sha}")
    lines.append("")
    lines.append("Global target mode weights (frozen-v4 logic):")
    lines.append(f"  line    = {morph['mode_weight_line']:.6f}")
    lines.append(f"  plane   = {morph['mode_weight_plane']:.6f}")
    lines.append(f"  compact = {morph['mode_weight_compact']:.6f}")
    lines.append("")
    lines.append("V5 parameters:")
    for k, val in config.items():
        lines.append(f"  {k} = {val}")
    lines.append("")
    lines.append("Pool sizes (filters unchanged from frozen v4):")
    for name, subset in pools.items():
        lines.append(f"  {name:>6}: {len(subset)}")
    lines.append("")

    for pool_name in ["all", "broad", "medium", "strict"]:
        ranked = sorted(
            pools[pool_name],
            key=lambda r: (-v4.safe_float(r.get("final_score_v5"), -1e99), str(r.get("pdb_name", ""))),
        )
        lines.append(f"Top {min(top_k, len(ranked))} in {pool_name} pool by final_score_v5:")
        for idx, r in enumerate(ranked[:top_k], start=1):
            lines.append(
                f"  #{idx:02d} {r.get('pdb_name')} "
                f"v5={v4.safe_float(r.get('final_score_v5')):.5f} "
                f"v4={v4.safe_float(r.get('final_score_v4')):.5f} "
                f"morph={v4.safe_float(r.get('morphology_blended_score_v5')):.5f} "
                f"extent={v4.safe_float(r.get('extent_score_v5')):.5f} "
                f"micro={v4.safe_float(r.get('score_microfit')):.5f} "
                f"level={r.get('filter_level')}"
            )
        lines.append("")

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


# -----------------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------------

def main() -> None:
    p = argparse.ArgumentParser(
        description="BinderRanker v5 prototype using SHA-guarded frozen v4 components."
    )
    p.add_argument("--input_csv", required=True, help="contact_field_v3 CSV")
    p.add_argument("--v4_script", required=True, help="path to frozen rank_candidates_v4_layered.py")
    p.add_argument("--output_xlsx", required=True)
    p.add_argument("--output_tsv", default=None)
    p.add_argument("--output_txt", default=None)
    p.add_argument("--top_n_sheets", type=int, default=500)
    p.add_argument("--top_k_report", type=int, default=30)

    # Keep v4 normalization/morphology controls explicit.
    p.add_argument("--morph_alpha", type=float, default=1.25)
    p.add_argument("--min_mode_weight", type=float, default=0.05)
    p.add_argument("--q_low", type=float, default=0.05)
    p.add_argument("--q_high", type=float, default=0.95)

    # V5 morphology architecture.  Defaults are v4-compatible (beta=0).
    p.add_argument("--beta_local", type=float, default=0.0,
                   help="blend global/local mode weights; 0 reproduces v4 morphology")
    p.add_argument("--local_lambda", type=float, default=1.0,
                   help="0=neutral local 1/3 weights, 1=full local morphology")
    p.add_argument("--gain_plane", type=float, default=1.0)
    p.add_argument("--gain_compact", type=float, default=1.0)

    # V5 extent architecture.
    p.add_argument(
        "--extent_variant",
        choices=["coverage", "coverage_cb", "coverage_micro", "coverage_continuity", "coverage_cb_micro"],
        default="coverage",
    )

    # V5 outer score. Defaults exactly match v4 outer composition.
    p.add_argument("--micro_weight", type=float, default=0.0)
    p.add_argument("--extent_weight", type=float, default=0.0)
    p.add_argument("--safety_weight", type=float, default=0.10)
    p.add_argument("--roughness_weight", type=float, default=0.05)

    args = p.parse_args()

    input_csv = Path(args.input_csv)
    v4_script = Path(args.v4_script)
    output_xlsx = Path(args.output_xlsx)
    output_tsv = Path(args.output_tsv) if args.output_tsv else output_xlsx.with_suffix(".scored.tsv")
    output_txt = Path(args.output_txt) if args.output_txt else output_xlsx.with_suffix(".txt")

    if output_xlsx.exists() or output_tsv.exists() or output_txt.exists():
        existing = [str(x) for x in [output_xlsx, output_tsv, output_txt] if x.exists()]
        raise SystemExit("[FATAL] refusing to overwrite existing outputs:\n  " + "\n  ".join(existing))

    v4, v4_sha = load_frozen_v4(v4_script)

    rows = v4.read_csv_rows(input_csv)
    if not rows:
        raise SystemExit("[FATAL] no rows found in input CSV")
    raw_original = [dict(r) for r in rows]

    # Frozen v4 calculation: source of all legacy component scores and filters.
    v4.add_derived_metrics(rows)
    morph = v4.infer_target_morphology(
        rows,
        alpha=args.morph_alpha,
        min_mode_weight=args.min_mode_weight,
    )
    v4_meta = v4.compute_scores(rows, morph, q_low=args.q_low, q_high=args.q_high)
    thresholds = v4.compute_filter_thresholds(rows)
    v4.apply_layered_filters(rows, thresholds)

    v5_config = add_v5_scores(
        rows=rows,
        morph=morph,
        v4=v4,
        beta_local=args.beta_local,
        local_lambda=args.local_lambda,
        gain_plane=args.gain_plane,
        gain_compact=args.gain_compact,
        morph_alpha=args.morph_alpha,
        min_mode_weight=args.min_mode_weight,
        q_low=args.q_low,
        q_high=args.q_high,
        extent_variant=args.extent_variant,
        micro_weight=args.micro_weight,
        extent_weight=args.extent_weight,
        safety_weight=args.safety_weight,
        roughness_weight=args.roughness_weight,
    )

    # Keep v4 ranks for side-by-side comparison and add v5 ranks separately.
    v4.add_ranks(rows, [
        "final_score_v4",
        "morphology_adaptive_score",
        "rank_consensus_score",
        "score_line",
        "score_plane",
        "score_compact",
        "score_roughness",
        "score_microfit",
        "score_safety",
    ])
    add_v5_ranks(rows, v4)

    headers = ordered_headers(rows)
    write_tsv(output_tsv, rows, headers)
    write_excel(
        output_xlsx,
        raw_original,
        rows,
        headers,
        v5_config,
        morph,
        v4_sha,
        args.top_n_sheets,
        v4,
    )
    write_report(
        output_txt,
        rows,
        v5_config,
        morph,
        v4_sha,
        args.top_k_report,
        v4,
    )

    # Compatibility audit: only meaningful when v5 args reduce to frozen v4.
    is_compat = (
        abs(args.beta_local) < EPS
        and abs(args.micro_weight) < EPS
        and abs(args.extent_weight) < EPS
        and abs(args.safety_weight - 0.10) < EPS
        and abs(args.roughness_weight - 0.05) < EPS
    )

    valid = [r for r in rows if v4.is_valid_row(r)]
    if is_compat:
        diffs = [
            abs(v4.safe_float(r.get("final_score_v5")) - v4.safe_float(r.get("final_score_v4")))
            for r in valid
        ]
        max_diff = max(diffs) if diffs else math.nan
        print(f"V4 compatibility max_abs_diff: {max_diff:.12g}")
        if math.isfinite(max_diff) and max_diff > 1e-10:
            raise SystemExit("[FATAL] v4 compatibility check failed")
        print("[PASS] V4-COMPATIBILITY CHECK")

    print("[PASS] BinderRanker v5 prototype ranking complete")
    print(f"Rows: {len(rows)}")
    print(f"Frozen v4 SHA256: {v4_sha}")
    print(f"Output XLSX: {output_xlsx}")
    print(f"Output TSV : {output_tsv}")
    print(f"Output TXT : {output_txt}")
    print("V5 config:")
    for k, val in v5_config.items():
        print(f"  {k} = {val}")


if __name__ == "__main__":
    main()
